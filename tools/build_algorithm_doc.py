from __future__ import annotations

from pathlib import Path
import re
import textwrap

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Aindra_Trader_Algorithm_Documentation.docx"


def read_source(path: str) -> list[str]:
    return (ROOT / path).read_text(encoding="utf-8").splitlines()


def extract_function(path: str, signature: str) -> str:
    lines = read_source(path)
    start = next(i for i, line in enumerate(lines) if signature in line)
    depth = 0
    end = start
    seen_open = False
    for i in range(start, len(lines)):
        line = lines[i]
        depth += line.count("{")
        if "{" in line:
            seen_open = True
        depth -= line.count("}")
        end = i
        if seen_open and depth <= 0:
            break
    return "\n".join(lines[start : end + 1])


def extract_lines(path: str, start: int, end: int) -> str:
    lines = read_source(path)
    return "\n".join(lines[start - 1 : end])


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_grid(table, widths: list[int]) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.append(grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)


def mark_header_row(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def shade_paragraph(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_title(doc: Document, title: str, subtitle: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run(title)
    run.font.name = "Calibri"
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor(31, 77, 120)

    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(10)
    r2 = p2.add_run(subtitle)
    r2.font.name = "Calibri"
    r2.font.size = Pt(11)
    r2.font.color.rgb = RGBColor(85, 85, 85)


def add_h1(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="Heading 1")
    p.add_run(text)


def add_h2(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="Heading 2")
    p.add_run(text)


def add_para(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(6)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_code(doc: Document, code: str, max_lines: int | None = None) -> None:
    lines = code.splitlines()
    if max_lines is not None and len(lines) > max_lines:
        lines = lines[:max_lines] + ["// ...snipped for readability; see referenced source file for full code"]
    wrapped: list[str] = []
    for line in lines:
        if len(line) <= 96:
            wrapped.append(line)
        else:
            indent = re.match(r"\s*", line).group(0)
            parts = textwrap.wrap(line, width=96, subsequent_indent=indent + "  ", replace_whitespace=False)
            wrapped.extend(parts or [line])

    for chunk_start in range(0, len(wrapped), 42):
        chunk = wrapped[chunk_start : chunk_start + 42]
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.0
        shade_paragraph(p, "F2F4F7")
        run = p.add_run("\n".join(chunk))
        run.font.name = "Courier New"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
        run.font.size = Pt(7.5)
        run.font.color.rgb = RGBColor(20, 30, 45)


def add_key_value_table(doc: Document, rows: list[tuple[str, str]], widths=(2700, 6660)) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_grid(table, list(widths))
    header = table.rows[0].cells
    header[0].text = "Item"
    header[1].text = "Current Value / Behavior"
    mark_header_row(table.rows[0])
    for cell in header:
        set_cell_shading(cell, "E8EEF5")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.bold = True
    for key, value in rows:
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = value
        for cell in cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_table_grid(table, list(widths))
    doc.add_paragraph()


def style_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, RGBColor(46, 116, 181), 18, 10),
        ("Heading 2", 13, RGBColor(46, 116, 181), 14, 7),
        ("Heading 3", 12, RGBColor(31, 77, 120), 10, 5),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ["List Bullet", "List Number"]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("Aindra Trader Algorithm Reference")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(85, 85, 85)


def build() -> None:
    doc = Document()
    style_doc(doc)
    add_title(
        doc,
        "Aindra Trader Algorithm Documentation",
        "Current paper-trading algorithm, risk guard, execution flow, and source-code excerpts.",
    )

    add_h1(doc, "1. Executive Summary")
    add_para(
        doc,
        "Aindra Trader currently uses a deterministic rule-based trading agent. It is not yet a trained LLM or reinforcement-learning model. The bot scans a fixed NSE cash watchlist, calculates technical signals, scores long and short setups, applies hard risk gates, paper-executes the highest eligible setup, and records trades in a backend ledger.",
    )
    add_bullets(
        doc,
        [
            "Trading mode today: paper orders. Live order code exists but remains locked unless LIVE_TRADING_ENABLED=true and Kite token validation passes.",
            "Instrument scope: NSE cash equities from the configured watchlist only; no futures or options.",
            "Risk model: fixed daily capital, fixed max daily loss, fixed risk per trade, and max trades per day.",
            "Exit model: daily target/loss, per-position stop/target, square-off window, market close, or manual close day.",
        ],
    )

    add_h1(doc, "2. Current Configuration")
    add_key_value_table(
        doc,
        [
            ("Capital", "Rs. 50,000"),
            ("Daily profit target", "1% = Rs. 500"),
            ("Daily max loss", "0.5% = Rs. 250"),
            ("Risk per trade", "0.25% = Rs. 125"),
            ("Max trades per day", "2 entries"),
            ("Minimum signal score", "70 / 100"),
            ("Maximum spread", "18 bps"),
            ("Stop loss distance", "0.8% minimum, ATR-adjusted when available"),
            ("Trade target distance", "1.6% minimum or 1.4x stop distance"),
            ("Slippage assumption", "8 bps per entry/exit side"),
            ("Strategy mode", "Hybrid: pick highest scoring setup"),
        ],
    )
    add_code(doc, extract_lines("src/App.tsx", 337, 349))
    add_code(doc, extract_lines("server/index.js", 32, 44))

    add_h1(doc, "3. Trade Universe")
    add_para(doc, "The scanner can select only symbols from the configured NSE cash watchlist. NIFTY 50 is used only as market context.")
    add_code(doc, extract_lines("server/watchlist.js", 1, 21))

    add_h1(doc, "4. Market Session Rules")
    add_para(doc, "The bot uses India Standard Time. It can be armed before the market opens, but fresh entries are only allowed after the early open buffer.")
    add_key_value_table(
        doc,
        [
            ("Pre-open", "09:00 to 09:15"),
            ("Market open", "09:15 to 15:29"),
            ("Fresh entries allowed", "09:20 to 15:00"),
            ("Square-off due", "From 15:15"),
            ("No-trade days", "Saturday, Sunday, and configured NSE holidays"),
        ],
    )
    add_code(doc, extract_function("server/marketCalendar.js", "export function getMarketSession"))

    add_h1(doc, "5. Indicators and Signal Inputs")
    add_para(doc, "For each quote, the strategy layer computes VWAP, ATR, volume pulse, RSI, Bollinger Bands, opening range, short-term momentum, day momentum, NIFTY bias, and spread bps.")
    add_code(doc, extract_lines("server/strategy.js", 1, 94), max_lines=95)

    add_h1(doc, "6. Strategy Scoring")
    add_para(doc, "The hybrid engine scores long and short versions of four strategy families: momentum, mean reversion, VWAP pullback, and opening range. The highest strategy score is selected when strategy mode is Hybrid.")
    add_bullets(
        doc,
        [
            "Long Momentum: rewards positive short-term/day momentum, volume pulse, positive NIFTY bias, and price above VWAP.",
            "Long Mean Reversion: rewards low RSI, proximity to lower Bollinger band, volume, and neutral market bias.",
            "Long VWAP Pullback: rewards price above VWAP, volume, positive day momentum, and controlled pullback pressure.",
            "Long Opening Range: rewards breakout above opening range high with volume and day momentum.",
            "Short variants mirror those conditions for downside setups.",
        ],
    )
    add_code(doc, extract_function("server/strategy.js", "function strategyScores"), max_lines=150)
    add_code(doc, extract_function("server/strategy.js", "function selectStrategy"))

    add_h1(doc, "7. Signal Eligibility")
    add_para(doc, "After selecting the best strategy for a symbol, the bot calculates a final score using spread quality, strategy quality, index alignment, and ATR risk quality.")
    add_key_value_table(
        doc,
        [
            ("Final score", "spreadScore + strategyScore + indexScore + riskScore, clamped to 0-100"),
            ("Score gate", "score >= minScore, currently 70"),
            ("Spread gate", "spreadBps <= maxSpreadBps, currently 18 bps"),
            ("Capital gate", "quote.ltp <= capital * 0.98"),
            ("Quantity gate", "risk sizing must produce at least 1 share"),
            ("Momentum gate", "directional momentum must pass, except mean-reversion setups"),
        ],
    )
    add_code(doc, extract_function("server/strategy.js", "export function generateSignals"), max_lines=130)

    add_h1(doc, "8. Agent Decision")
    add_para(doc, "The agent picks the first eligible signal after sorting by score. If no signal is eligible, it waits and reports why.")
    add_code(doc, extract_function("server/strategy.js", "export function agentDecision"))

    add_h1(doc, "9. Position Sizing")
    add_para(doc, "The bot sizes positions using both a capital cap and a risk cap. Quantity is the lower of the two.")
    add_code(
        doc,
        """riskAmount = capital * riskPerTradePct / 100
percentStopDistance = price * stopLossPct / 100
atrStopDistance = atrValue * 0.7
stopDistance = max(percentStopDistance, atrStopDistance, price * 0.004)
capitalCapQuantity = floor(capital * 0.96 / price)
riskQuantity = floor(riskAmount / stopDistance)
quantity = min(capitalCapQuantity, riskQuantity)
targetDistance = max(price * takeProfitPct / 100, stopDistance * 1.4)""",
    )
    add_code(doc, extract_function("server/strategy.js", "export function sizePosition"))

    add_h1(doc, "10. Pricing, Slippage, P&L, and Charges")
    add_para(doc, "The paper engine applies an 8 bps slippage assumption and deducts estimated equity intraday charges from realized and unrealized P&L.")
    add_key_value_table(
        doc,
        [
            ("Long entry", "price * (1 + slippageBps / 10000)"),
            ("Long exit", "price * (1 - slippageBps / 10000)"),
            ("Short entry", "price * (1 - slippageBps / 10000)"),
            ("Short exit", "price * (1 + slippageBps / 10000)"),
        ],
    )
    add_code(doc, extract_function("src/App.tsx", "function calcCharges"))
    add_code(doc, extract_function("src/App.tsx", "function estimatePositionPnl"))
    add_code(doc, extract_lines("src/App.tsx", 745, 753))

    add_h1(doc, "11. Runtime Trading Loop")
    add_para(doc, "The frontend paper engine polls backend strategy/risk payloads, updates local positions, exits first when limits are hit, and only opens a new position when no position is currently open and the daily entry count allows it.")
    add_bullets(
        doc,
        [
            "If market is closed and the bot cannot remain armed, close open paper positions and stop.",
            "If daily target or daily loss is reached, close all open positions and stop.",
            "If square-off is due, close open positions.",
            "If a position stop or target is hit, close that position.",
            "If max trades are already taken, block new entries.",
            "If no position exists and strategy decision is TRADE, open the recommended long or short position.",
        ],
    )
    add_code(doc, extract_function("src/App.tsx", "function tickBackendPaper"), max_lines=170)

    add_h1(doc, "12. Risk Guard")
    add_para(doc, "Every order passes a server-side risk validator before being accepted into the paper ledger or sent to live order code.")
    add_key_value_table(
        doc,
        [
            ("Kill switch", "Reject order when active"),
            ("Market phase", "Reject fresh entries outside market/open-entry window"),
            ("Exchange/product", "Only NSE and MIS are allowed"),
            ("Quantity", "Must be positive"),
            ("Capital", "Entry notional must be <= 98% of configured capital"),
            ("Max trades", "Reject new entries once max trades reached"),
            ("Daily loss", "Reject new entries after max daily loss is reached"),
            ("Live-only checks", "Kite token must exist and LIVE_TRADING_ENABLED must be true"),
        ],
    )
    add_code(doc, extract_function("server/riskGuard.js", "export function validatePaperOrder"))
    add_code(doc, extract_function("server/riskGuard.js", "export function validateOrder"), max_lines=80)

    add_h1(doc, "13. Order Recording and Trade History")
    add_para(doc, "Paper orders are stored as events, deduplicated by clientOrderId, and then reconstructed into trade history by matching ENTRY and EXIT lots.")
    add_code(doc, extract_function("server/orders.js", "export function paperOrder"), max_lines=90)
    add_code(doc, extract_function("server/orders.js", "export function tradeHistory"), max_lines=150)

    add_h1(doc, "14. Current Limitations and Verification Notes")
    add_bullets(
        doc,
        [
            "This is a rule-based signal engine, not a trained best-trader AI model.",
            "It trades only NSE cash equity symbols in the watchlist; futures/options are not included.",
            "Paper fills use the quoted/estimated price plus configured slippage; real fills may differ.",
            "Historical performance must be verified with longer backtests, forward paper trading, and then very small live tests.",
            "Live orders remain locked until explicitly enabled and verified with Zerodha token/session checks.",
            "Any unmatched old paper entries should remain reset/archived unless manually repaired with exact exit price/time.",
        ],
    )

    add_h1(doc, "15. Source File Map")
    add_key_value_table(
        doc,
        [
            ("server/strategy.js", "Indicators, strategy scoring, signal generation, and agent decision."),
            ("server/riskGuard.js", "Daily risk state, kill switch, market-entry validation, live/paper order gates."),
            ("server/marketCalendar.js", "IST market calendar, entry windows, square-off timing, next open time."),
            ("server/orders.js", "Paper/live order recording, dedupe, trade-history reconstruction."),
            ("server/watchlist.js", "Allowed NSE symbols and NIFTY context symbol."),
            ("src/App.tsx", "Frontend runtime loop, local paper positions, P&L, slippage, UI controls, reports."),
            ("server/index.js", "API routes and server-side default configuration."),
        ],
    )

    doc.core_properties.title = "Aindra Trader Algorithm Documentation"
    doc.core_properties.subject = "Trading algorithm and code reference"
    doc.core_properties.author = "Codex"
    doc.save(OUT)


if __name__ == "__main__":
    build()
    print("Aindra_Trader_Algorithm_Documentation.docx")
